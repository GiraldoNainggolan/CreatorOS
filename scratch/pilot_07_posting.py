import os
import re
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Flowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def reconstruct_text(file_path):
    with open(file_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    cleaned_lines = [line.strip() for line in lines if line.strip()]
    
    reconstructed = []
    current_paragraph = []
    
    structural_markers = [
        r'^\d{2}_[A-Z_]+', # 01_PUBLISHING_STRATEGY
        r'^📂',
        r'^🏛',
        r'^##',
        r'^###',
        r'^\[ \]',
        r'^●',
        r'^\-',
        r'^\*',
        r'^\d+\.', # 1., 2.
        r'^====',
        r'^\[[A-Z\s&_]+\]', # [CONTENT DISTRIBUTION MET ADA TA]
        r'^Plaintext'
    ]
    
    for word in cleaned_lines:
        is_marker = False
        for marker in structural_markers:
            if re.match(marker, word):
                is_marker = True
                break
        
        # Specific fix for common split words, ONLY if we have very high confidence.
        # But user said: DO NOT aggressively use regex. "boleh diperbaiki HANYA jika confidence sangat tinggi"
        # We will just append the word for now.
        
        if is_marker:
            if current_paragraph:
                reconstructed.append(" ".join(current_paragraph))
                current_paragraph = []
            reconstructed.append(word)
        else:
            current_paragraph.append(word)
            
    if current_paragraph:
        reconstructed.append(" ".join(current_paragraph))
        
    full_text = "\n\n".join(reconstructed)
    
    # Post-processing to fix metadata blocks that got squashed
    # E.g. "Content ID : PUB-2026" -> insert newline before "Content ID"
    # Looking for a pattern of Capitalized words followed by colon.
    full_text = re.sub(r'\s+([A-Z][a-zA-Z\s]+) :', r'\n\1 :', full_text)
    
    # Also fix some split words we know are just artifacts (high confidence)
    full_text = full_text.replace("STRA TEGY", "STRATEGY")
    full_text = full_text.replace("ANAL YTICS", "ANALYTICS")
    full_text = full_text.replace("MET ADA TA", "METADATA")
    full_text = full_text.replace("DIREKT ORI", "DIREKTORI")
    full_text = full_text.replace("PLA TFORM", "PLATFORM")
    full_text = full_text.replace("AUT OMA TION", "AUTOMATION")
    full_text = full_text.replace("TEMPLA TE", "TEMPLATE")
    full_text = full_text.replace("Contr ol", "Control")
    
    return full_text

def wrap_with_headers(domain, content, source_pages="219-223"):
    header = f"""# CREATOROS — {domain}

## Source Information
Source: Paket_Lengkap.pdf
Source section: {domain}
Source pages: {source_pages}

# ORIGINAL SOURCE STRUCTURE

"""
    footer = f"""

# SOURCE TRACEABILITY
- Source file: knowledge/Paket_Lengkap.pdf
- Page range: {source_pages}
- Extraction file: pdf_extracts/{domain}.txt
- Generation date: 2026-08-09
- Content status: CANONICAL RECONSTRUCTED
"""
    return header + content + footer

def generate_docx(md_content, out_path):
    doc = Document()
    for line in md_content.split("\n"):
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.strip() == "":
            continue
        else:
            doc.add_paragraph(line)
    doc.save(out_path)

def generate_pdf(md_content, out_path):
    doc = SimpleDocTemplate(out_path, pagesize=letter)
    styles = getSampleStyleSheet()
    Story: list[Flowable] = []
    
    for line in md_content.split("\n"):
        if line.startswith("# "):
            p = Paragraph(f"<b>{line[2:]}</b>", styles['Heading1'])
        elif line.startswith("## "):
            p = Paragraph(f"<b>{line[3:]}</b>", styles['Heading2'])
        elif line.startswith("### "):
            p = Paragraph(f"<b>{line[4:]}</b>", styles['Heading3'])
        elif line.strip() == "":
            p = Spacer(1, 12)
        else:
            p = Paragraph(line, styles['Normal'])
        Story.append(p)
        Story.append(Spacer(1, 6))
        
    doc.build(Story)

def main():
    domain = "07_POSTING"
    source_txt = rf"c:\Users\Giraldo Nainggolan\Desktop\Clude\ECC\project\CreatorOS\scratch\pdf_extracts\{domain}.txt"
    out_dir = rf"c:\Users\Giraldo Nainggolan\Desktop\Clude\ECC\project\CreatorOS\knowledge\{domain}"
    
    os.makedirs(out_dir, exist_ok=True)
    
    md_path = os.path.join(out_dir, f"{domain}.md")
    docx_path = os.path.join(out_dir, f"{domain}.docx")
    pdf_path = os.path.join(out_dir, f"{domain}.pdf")
    
    # 1. Reconstruct text
    raw_reconstructed = reconstruct_text(source_txt)
    final_md_content = wrap_with_headers(domain, raw_reconstructed)
    
    # 2. Write MD
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(final_md_content)
        
    # 3. Write DOCX
    generate_docx(final_md_content, docx_path)
    
    # 4. Write PDF
    generate_pdf(final_md_content, pdf_path)
    
    print(f"PILOT GENERATION COMPLETE FOR {domain}")
    print(f"MD, DOCX, PDF saved to {out_dir}")

if __name__ == "__main__":
    main()
